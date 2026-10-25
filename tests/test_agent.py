"""Offline tests: a fake LLM replaces Groq, so these run without an API key or internet."""
import copy

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import config, llm, main
from app.errors import LLMConfigError, LLMError

REQUEST = {"request": "Create a project plan for launching a support chatbot over 3 months."}

PLAN = {
    "doc_type": "project_plan",
    "title": "Support Chatbot Launch Plan",
    "audience": "Executive sponsors",
    "assumptions": ["Team of 5 is available", "English only at launch"],
    "conflicts": [{"issue": "3 months vs full feature set", "resolution": "Phase 2 for advanced features"}],
    "tasks": [
        {"id": 7, "title": "Budget data", "tool": "get_mock_data", "data_type": "budget", "description": "x"},
        {"id": 9, "title": "Timeline data", "tool": "get_mock_data", "data_type": "timeline", "description": "x"},
        {"id": 3, "title": "Overview", "tool": "write_section", "data_type": None, "description": "x"},
        {"id": 4, "title": "Schedule", "tool": "build_table", "data_type": None, "description": "x"},
        {"id": 5, "title": "Risks", "tool": "write_section", "data_type": None, "description": "x"},
    ],
}
PROSE = {"heading": "Overview", "paragraphs": ["Some paragraph."], "bullets": ["one", "two"]}
TABLE = {"heading": "Schedule", "intro": "Phases below.", "columns": ["Phase", "Weeks"], "rows": [["Build", 5], ["Test"]]}
SUMMARY = {"heading": "ignored", "paragraphs": ["Summary text."], "bullets": ["Key point"]}
REVIEW = {
    "verdict": "revise",
    "issues": [
        {"section_index": 0, "problem": "Too vague", "fix": "Add specifics"},
        {"section_index": None, "problem": "No communication plan", "fix": "Add one"},
    ],
}


def make_fake(review=None, fail_labels=(), plans=None):
    calls = {"plan": 0}

    def fake(system, user, *, max_tokens=1500, label="llm"):
        if label in fail_labels:
            raise LLMError(f"forced failure in {label}")
        if label == "plan":
            calls["plan"] += 1
            return copy.deepcopy(plans[min(calls["plan"], len(plans)) - 1] if plans else PLAN)
        if label in ("write_section", "revise_section", "add_section"):
            return copy.deepcopy(PROSE) if label != "add_section" else {**PROSE, "heading": "Communication Plan"}
        if label == "build_table":
            return copy.deepcopy(TABLE)
        if label == "summary":
            return copy.deepcopy(SUMMARY)
        if label == "review":
            return copy.deepcopy(review if review is not None else {"verdict": "pass", "issues": []})
        raise AssertionError(f"unexpected label {label}")

    fake.calls = calls
    return fake


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path)
    return TestClient(main.app)


# ---------------- API behaviour ----------------
def test_full_flow_with_reflection_repairs(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake(review=REVIEW))
    res = client.post("/agent", json=REQUEST)
    assert res.status_code == 200
    body = res.json()

    assert body["status"] == "completed"
    assert [t["id"] for t in body["plan"]["tasks"]] == [1, 2, 3, 4, 5]  # ids renumbered by code
    assert all(t["status"] == "done" for t in body["plan"]["tasks"])
    assert body["reflection"]["verdict"] == "revised"
    assert body["reflection"]["sections_revised"] == ["Overview"]
    assert body["reflection"]["sections_added"] == ["Communication Plan"]

    download = client.get(body["download_url"])
    assert download.status_code == 200
    path = next(config.OUTPUT_DIR.glob("*.docx"))
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "Support Chatbot Launch Plan" in text
    assert "Executive Summary" in text
    assert "Communication Plan" in text
    table = Document(path).tables[-1]  # the Schedule table; short row was padded
    assert [c.text for c in table.rows[2].cells] == ["Test", ""]


def test_reflection_pass_changes_nothing(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake())
    body = client.post("/agent", json=REQUEST).json()
    assert body["reflection"]["verdict"] == "pass"
    assert body["reflection"]["sections_revised"] == []


@pytest.mark.parametrize("payload", [{}, {"request": ""}, {"request": "short"}, {"request": "x" * 2001}])
def test_request_validation(client, payload):
    assert client.post("/agent", json=payload).status_code == 422


def test_failed_task_does_not_kill_the_run(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake(fail_labels=("build_table",)))
    body = client.post("/agent", json=REQUEST).json()
    assert body["status"] == "completed_with_errors"
    failed = [t for t in body["plan"]["tasks"] if t["status"] == "failed"]
    assert len(failed) == 1 and failed[0]["tool"] == "build_table" and failed[0]["error"]
    assert client.get(body["download_url"]).status_code == 200


def test_reflection_failure_is_not_fatal(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake(fail_labels=("review",)))
    body = client.post("/agent", json=REQUEST).json()
    assert body["status"] == "completed"
    assert body["reflection"]["verdict"] == "skipped"


def test_summary_failure_uses_fallback(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake(fail_labels=("summary",)))
    assert client.post("/agent", json=REQUEST).status_code == 200


def test_planner_retries_after_invalid_plan(client, monkeypatch):
    bad = copy.deepcopy(PLAN)
    bad["tasks"][2]["tool"] = "send_email"  # not a real tool
    fake = make_fake(plans=[bad, PLAN])
    monkeypatch.setattr(llm, "chat_json", fake)
    assert client.post("/agent", json=REQUEST).status_code == 200
    assert fake.calls["plan"] == 2


def test_planner_gives_up_with_502(client, monkeypatch):
    bad = copy.deepcopy(PLAN)
    bad["tasks"] = bad["tasks"][:1]
    monkeypatch.setattr(llm, "chat_json", make_fake(plans=[bad, bad]))
    res = client.post("/agent", json=REQUEST)
    assert res.status_code == 502


def test_all_content_tasks_failing_returns_502(client, monkeypatch):
    monkeypatch.setattr(llm, "chat_json", make_fake(fail_labels=("write_section", "build_table")))
    assert client.post("/agent", json=REQUEST).status_code == 502


def test_missing_api_key_returns_503(client, monkeypatch):
    def no_key(*args, **kwargs):
        raise LLMConfigError("GROQ_API_KEY is missing")

    monkeypatch.setattr(llm, "chat_json", no_key)
    res = client.post("/agent", json=REQUEST)
    assert res.status_code == 503 and "GROQ_API_KEY" in res.json()["detail"]


@pytest.mark.parametrize("bad_id", ["not-an-id", "..%2F..%2F.env", "0" * 32])
def test_download_rejects_bad_ids(client, bad_id):
    assert client.get(f"/download/{bad_id}").status_code == 404


# ---------------- LLM layer ----------------
def test_parse_json_handles_fences_and_chatter():
    assert llm.parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert llm.parse_json('Sure! Here you go: {"a": {"b": 2}} Hope it helps.') == {"a": {"b": 2}}
    with pytest.raises(ValueError):
        llm.parse_json("no json here")
    with pytest.raises(ValueError):
        llm.parse_json("[1, 2, 3]")


def test_chat_json_retries_then_falls_back(monkeypatch):
    monkeypatch.setattr(llm.time, "sleep", lambda s: None)
    monkeypatch.setattr(config, "PRIMARY_MODEL", "big")
    monkeypatch.setattr(config, "FALLBACK_MODEL", "small")
    seen = []

    def flaky(model, system, user, max_tokens):
        seen.append(model)
        if model == "big":
            raise RuntimeError("429 rate limited")
        return '```json\n{"ok": true}\n```'

    monkeypatch.setattr(llm, "_call", flaky)
    assert llm.chat_json("s", "u") == {"ok": True}
    assert seen == ["big", "big", "small"]


def test_chat_json_raises_when_everything_fails(monkeypatch):
    monkeypatch.setattr(llm.time, "sleep", lambda s: None)
    monkeypatch.setattr(llm, "_call", lambda *a: "garbage")
    with pytest.raises(LLMError):
        llm.chat_json("s", "u", label="t")


def test_missing_key_is_not_retried(monkeypatch):
    monkeypatch.setattr(config, "GROQ_API_KEY", "")
    monkeypatch.setattr(llm, "_client", None)
    with pytest.raises(LLMConfigError):
        llm.chat_json("s", "u")


# ---------------- schema leniency (small LLMs are sloppy) ----------------
def test_plan_accepts_sloppy_llm_values():
    from app.schemas import Plan

    sloppy = copy.deepcopy(PLAN)
    sloppy["tasks"][0]["id"] = "task-one"
    sloppy["tasks"][2]["data_type"] = "none"
    sloppy["tasks"][3]["tool"] = " Build_Table "
    sloppy["conflicts"] = ["Only a string conflict"]
    plan = Plan(**sloppy)
    assert plan.tasks[2].data_type is None
    assert plan.tasks[3].tool == "build_table"
    assert plan.conflicts[0].issue == "Only a string conflict"
